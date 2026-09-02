import csv
import hashlib
import io
import re
import uuid
import zipfile
from datetime import date, datetime, timedelta, timezone
from types import SimpleNamespace
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select, update

import app.payouts.policy as payout_policy
import app.payouts.service as payout_service
import app.exports.payout_data as export_payout_data
from app.api.v1.exports import get_artifact_store
from app.audit.models import SecurityEvent
from app.audit.service import AuditAction, AuditContext
from app.auth.models import AccountStatus, Role, User
from app.auth.security import hash_password
from app.billing.models import BalanceLedgerEntry, CreatorBalance
from app.creators.models import CreatorProfile, ProfileStatus, RecipientStatus
from app.database.session import get_db
from app.contracts import ExportCommand
from app.exports.processor import execute_export
from app.exports.storage import ArtifactDownload
from app.exports.generator import moscow_date
from app.payouts.models import (
    ANONYMIZED_TEXT,
    PayoutDetails,
    PayoutEvent,
    PayoutEventDetail,
    PayoutRequest,
)
from app.payouts.policy import MOSCOW, moscow_today
from app.payouts.retention import anonymize_eligible_payout_pii
from app.payouts.schemas import PayoutReviewRequest
from app.scheduling.exports import dispatch_export


PASSWORD = "payout-test-password-123"
VALID_SBP_PHONE = "+79991234567"
COMMAND_RECEIPT_FIELDS = {
    "payout_request_id",
    "request_number",
    "sequence_number",
    "action",
    "from_status",
    "status",
    "actor_user_id",
    "recorded_at",
    "details",
}
MY_PAYOUT_FORBIDDEN_FIELDS = {
    "blogger_id",
    "manager_comment",
    "reviewer_user_id",
    "requisites_verified_at",
    "requisites_verified_by_user_id",
    "self_employment_verified_at",
    "self_employment_verified_by_user_id",
    "approved_by_user_id",
    "paid_recorded_at",
    "paid_by_user_id",
    "payment_reference",
    "rejected_by_user_id",
    "receipt_recorded_at",
    "receipt_received_by_user_id",
}


def _csrf_headers(client) -> dict[str, str]:
    return {"X-CSRF-Token": client.cookies.get("amp_csrf")}


def _assert_private_no_store(response) -> None:
    assert response.headers["cache-control"] == "private, no-store"


def _login(client, email: str) -> None:
    response = client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": PASSWORD},
    )
    assert response.status_code == 200, response.text


def _seed_staff(client, role: Role, label: str) -> tuple[uuid.UUID, str]:
    email = f"{label}-{uuid.uuid4()}@example.com"
    with client.app.state.test_session() as db:
        user = User(
            email=email,
            password_hash=hash_password(PASSWORD),
            role=role,
            status=AccountStatus.ACTIVE,
            email_verified_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
        )
        db.add(user)
        db.commit()
        return user.id, email


def _seed_creator(
    client,
    *,
    label: str,
    available_kopecks: int,
    recipient_status: RecipientStatus = RecipientStatus.INDIVIDUAL,
    full_name: str = "Test Creator",
    display_name: str = "Test Channel",
) -> tuple[uuid.UUID, str]:
    email = f"{label}-{uuid.uuid4()}@example.com"
    with client.app.state.test_session() as db:
        user = User(
            email=email,
            password_hash=hash_password(PASSWORD),
            role=Role.BLOGGER,
            status=AccountStatus.ACTIVE,
            email_verified_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
        )
        db.add(user)
        db.flush()
        db.add_all(
            [
                CreatorProfile(
                    user_id=user.id,
                    full_name=full_name,
                    display_name=display_name,
                    phone="+79990000000",
                    telegram="@payout_test",
                    city_country="Moscow, Russia",
                    content_topics="Technology",
                    recipient_status=recipient_status,
                    status=ProfileStatus.APPROVED,
                    submitted_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
                    reviewed_at=datetime(2026, 1, 2, tzinfo=timezone.utc),
                ),
                CreatorBalance(
                    blogger_id=user.id,
                    available_kopecks=available_kopecks,
                    reserved_kopecks=0,
                    paid_kopecks=0,
                ),
            ]
        )
        db.commit()
        return user.id, email


def _put_details(client, *, phone: str = VALID_SBP_PHONE, bank_name: str | None = None):
    return client.put(
        "/api/v1/me/payout-details",
        json={"sbp_phone": phone, "bank_name": bank_name},
        headers=_csrf_headers(client),
    )


def _create_request(client, *, idempotency_key: uuid.UUID | None = None):
    return client.post(
        "/api/v1/me/payout-requests",
        json={"idempotency_key": str(idempotency_key or uuid.uuid4())},
        headers=_csrf_headers(client),
    )


def _staff_command(client, payout_id: str, command: str, payload: dict):
    return client.post(
        f"/api/v1/staff/payout-requests/{payout_id}/{command}",
        json=payload,
        headers=_csrf_headers(client),
    )


def _review_and_approve(
    client,
    *,
    payout_id: str,
    manager_email: str,
    self_employed: bool,
) -> dict:
    _login(client, manager_email)
    reviewed = _staff_command(
        client,
        payout_id,
        "reviews",
        {"idempotency_key": str(uuid.uuid4())},
    )
    assert reviewed.status_code == 200, reviewed.text
    approved = _staff_command(
        client,
        payout_id,
        "approvals",
        {
            "idempotency_key": str(uuid.uuid4()),
            "requisites_verified": True,
            "self_employment_verified": True if self_employed else None,
        },
    )
    assert approved.status_code == 200, approved.text
    return approved.json()


def _prepare_request(
    client,
    *,
    label: str,
    available_kopecks: int,
    recipient_status: RecipientStatus = RecipientStatus.INDIVIDUAL,
    full_name: str = "Test Creator",
    display_name: str = "Test Channel",
    bank_name: str | None = "Test Bank",
) -> tuple[uuid.UUID, str, dict]:
    blogger_id, email = _seed_creator(
        client,
        label=label,
        available_kopecks=available_kopecks,
        recipient_status=recipient_status,
        full_name=full_name,
        display_name=display_name,
    )
    _login(client, email)
    details = _put_details(client, bank_name=bank_name)
    assert details.status_code == 200, details.text
    created = _create_request(client)
    assert created.status_code == 201, created.text
    receipt = created.json()
    payout = client.get(
        f"/api/v1/me/payout-requests/{receipt['payout_request_id']}"
    )
    assert payout.status_code == 200, payout.text
    return blogger_id, email, payout.json()


def _balance(client, blogger_id: uuid.UUID) -> CreatorBalance:
    with client.app.state.test_session() as db:
        balance = db.get(CreatorBalance, blogger_id)
        assert balance is not None
        db.expunge(balance)
        return balance


def _ledger(client, blogger_id: uuid.UUID) -> list[BalanceLedgerEntry]:
    with client.app.state.test_session() as db:
        entries = list(
            db.scalars(
                select(BalanceLedgerEntry)
                .where(BalanceLedgerEntry.blogger_id == blogger_id)
                .order_by(BalanceLedgerEntry.created_at, BalanceLedgerEntry.id)
            )
        )
        for entry in entries:
            db.expunge(entry)
        return entries


def _enum_value(value):
    return getattr(value, "value", value)


def test_command_hash_ignores_new_unset_optional_schema_fields():
    class EvolvedPayoutReviewRequest(PayoutReviewRequest):
        future_optional_field: str | None = None

    actor_id = uuid.uuid4()
    payout_id = uuid.uuid4()
    idempotency_key = uuid.uuid4()
    old_payload = PayoutReviewRequest(
        idempotency_key=idempotency_key,
        comment="Stable command",
    )
    evolved_payload = EvolvedPayoutReviewRequest(
        idempotency_key=idempotency_key,
        comment="Stable command",
    )

    assert payout_service._command_hash(
        action="review_started",
        actor_id=actor_id,
        payout_request_id=payout_id,
        payload=old_payload,
    ) == payout_service._command_hash(
        action="review_started",
        actor_id=actor_id,
        payout_request_id=payout_id,
        payload=evolved_payload,
    )


def test_payout_details_require_csrf_validate_phone_and_remain_owner_scoped(client):
    first_id, first_email = _seed_creator(
        client,
        label="details-first",
        available_kopecks=100,
    )
    _, second_email = _seed_creator(
        client,
        label="details-second",
        available_kopecks=100,
    )

    _login(client, first_email)
    assert client.put(
        "/api/v1/me/payout-details",
        json={"sbp_phone": VALID_SBP_PHONE, "bank_name": "First Bank"},
    ).status_code == 403
    invalid = _put_details(client, phone="8999-123")
    assert invalid.status_code == 422

    saved = _put_details(client, bank_name="First Bank")
    assert saved.status_code == 200, saved.text
    assert saved.json()["blogger_id"] == str(first_id)
    assert saved.json()["sbp_phone"] == VALID_SBP_PHONE
    assert saved.json()["bank_name"] == "First Bank"

    _login(client, second_email)
    second = _put_details(client, phone="+79997654321", bank_name="Second Bank")
    assert second.status_code == 200, second.text
    assert client.get("/api/v1/me/payout-details").json()["bank_name"] == "Second Bank"

    _login(client, first_email)
    mine = client.get("/api/v1/me/payout-details")
    assert mine.status_code == 200
    assert mine.json()["blogger_id"] == str(first_id)
    assert mine.json()["bank_name"] == "First Bank"


def test_suspended_blogger_can_read_but_cannot_change_payout_data(client):
    blogger_id, email, payout = _prepare_request(
        client,
        label="suspended-reader",
        available_kopecks=12_345,
    )
    with client.app.state.test_session() as db:
        blogger = db.get(User, blogger_id)
        blogger.status = AccountStatus.SUSPENDED
        blogger.status_changed_at = datetime.now(timezone.utc)
        db.commit()

    _login(client, email)
    details = client.get("/api/v1/me/payout-details")
    listing = client.get("/api/v1/me/payout-requests")
    detail = client.get(f"/api/v1/me/payout-requests/{payout['id']}")
    assert details.status_code == 200, details.text
    assert listing.status_code == 200, listing.text
    assert detail.status_code == 200, detail.text

    assert _put_details(client, bank_name="Changed Bank").status_code == 403
    assert _create_request(client).status_code == 403


def test_create_reserves_entire_balance_once_and_has_one_active_request(client):
    blogger_id, email = _seed_creator(
        client,
        label="reserve",
        available_kopecks=123_456,
    )
    _login(client, email)
    assert _put_details(client).status_code == 200

    key = uuid.uuid4()
    without_csrf = client.post(
        "/api/v1/me/payout-requests",
        json={"idempotency_key": str(key)},
    )
    assert without_csrf.status_code == 403

    created = _create_request(client, idempotency_key=key)
    assert created.status_code == 201, created.text
    _assert_private_no_store(created)
    body = created.json()
    assert set(body) == COMMAND_RECEIPT_FIELDS
    assert body["sequence_number"] == 1
    assert body["action"] == "requested"
    assert body["from_status"] is None
    assert body["status"] == "requested"
    assert body["actor_user_id"] == str(blogger_id)
    assert body["details"] == {"amount_kopecks": 123_456}
    assert "command_schema_version" not in body["details"]

    payout_id = body["payout_request_id"]
    mine = client.get(f"/api/v1/me/payout-requests/{payout_id}")
    assert mine.status_code == 200, mine.text
    _assert_private_no_store(mine)
    assert mine.json()["amount_kopecks"] == 123_456
    assert mine.json()["currency"] == "RUB"
    assert not (MY_PAYOUT_FORBIDDEN_FIELDS & mine.json().keys())

    balance = _balance(client, blogger_id)
    assert (balance.available_kopecks, balance.reserved_kopecks, balance.paid_kopecks) == (
        0,
        123_456,
        0,
    )
    entries = _ledger(client, blogger_id)
    assert len(entries) == 1
    assert _enum_value(entries[0].operation_type) == "payout_reserved"
    assert (
        entries[0].available_delta_kopecks,
        entries[0].reserved_delta_kopecks,
        entries[0].paid_delta_kopecks,
    ) == (-123_456, 123_456, 0)
    assert entries[0].idempotency_key == f"payout-reserve:{payout_id}"

    replay = _create_request(client, idempotency_key=key)
    assert replay.status_code == 201, replay.text
    assert replay.json() == body
    assert len(_ledger(client, blogger_id)) == 1

    second = _create_request(client)
    assert second.status_code == 409
    assert second.json()["error"]["code"] == "PAYOUT_ALREADY_ACTIVE"

    with client.app.state.test_session() as db:
        assert db.scalar(
            select(func.count()).select_from(PayoutRequest).where(
                PayoutRequest.blogger_id == blogger_id
            )
        ) == 1
        event = db.scalar(
            select(PayoutEvent).where(
                PayoutEvent.payout_request_id == uuid.UUID(payout_id)
            )
        )
        assert event.event_metadata["command_schema_version"] == 1


def test_request_rolls_back_balance_ledger_history_and_audit_together(
    client,
    monkeypatch,
):
    blogger_id, email = _seed_creator(
        client,
        label="atomic",
        available_kopecks=50_000,
    )
    _login(client, email)
    assert _put_details(client).status_code == 200

    def fail_audit(*args, **kwargs):
        raise RuntimeError("audit unavailable")

    monkeypatch.setattr(payout_service, "record_event", fail_audit)
    with pytest.raises(RuntimeError, match="audit unavailable"):
        _create_request(client)

    balance = _balance(client, blogger_id)
    assert (balance.available_kopecks, balance.reserved_kopecks, balance.paid_kopecks) == (
        50_000,
        0,
        0,
    )
    with client.app.state.test_session() as db:
        assert db.scalar(select(func.count()).select_from(PayoutRequest)) == 0
        assert db.scalar(select(func.count()).select_from(PayoutEvent)) == 0
        assert db.scalar(select(func.count()).select_from(BalanceLedgerEntry)) == 0
        assert db.scalar(
            select(func.count())
            .select_from(SecurityEvent)
            .where(SecurityEvent.object_type == "payout_request")
        ) == 0


def test_http_commit_failure_never_returns_success_and_rolls_back_command(client):
    blogger_id, email = _seed_creator(
        client,
        label="commit-failure",
        available_kopecks=54_300,
    )
    _login(client, email)
    assert _put_details(client).status_code == 200

    session_factory = client.app.state.test_session
    resolved_sessions = []

    def failing_get_db():
        db = session_factory()
        resolved_sessions.append(db)

        def fail_commit():
            raise RuntimeError("forced payout commit failure")

        db.commit = fail_commit
        try:
            yield db
            db.commit()
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    original_override = client.app.dependency_overrides[get_db]
    transport = client._transport
    original_raise_server_exceptions = transport.raise_server_exceptions
    client.app.dependency_overrides[get_db] = failing_get_db
    transport.raise_server_exceptions = False
    try:
        response = _create_request(client)
    finally:
        transport.raise_server_exceptions = original_raise_server_exceptions
        client.app.dependency_overrides[get_db] = original_override

    assert response.status_code == 500, response.text
    assert len(resolved_sessions) == 1

    balance = _balance(client, blogger_id)
    assert (balance.available_kopecks, balance.reserved_kopecks, balance.paid_kopecks) == (
        54_300,
        0,
        0,
    )
    with session_factory() as db:
        assert db.scalar(select(func.count()).select_from(PayoutRequest)) == 0
        assert db.scalar(select(func.count()).select_from(PayoutEvent)) == 0
        assert db.scalar(select(func.count()).select_from(BalanceLedgerEntry)) == 0
        assert db.scalar(
            select(func.count()).select_from(SecurityEvent).where(
                SecurityEvent.object_type == "payout_request"
            )
        ) == 0
    _assert_private_no_store(response)


def test_self_employed_approval_enforces_state_rbac_and_verifications(client):
    blogger_id, blogger_email, payout = _prepare_request(
        client,
        label="approval",
        available_kopecks=90_000,
        recipient_status=RecipientStatus.SELF_EMPLOYED,
    )
    manager_id, manager_email = _seed_staff(client, Role.MANAGER, "approval-manager")
    _, finance_email = _seed_staff(client, Role.FINANCE, "approval-finance")

    assert client.get("/api/v1/staff/payout-requests").status_code == 403
    assert _staff_command(
        client,
        payout["id"],
        "reviews",
        {"idempotency_key": str(uuid.uuid4())},
    ).status_code == 403

    _login(client, manager_email)
    before_review = _staff_command(
        client,
        payout["id"],
        "approvals",
        {
            "idempotency_key": str(uuid.uuid4()),
            "requisites_verified": True,
            "self_employment_verified": True,
        },
    )
    assert before_review.status_code == 409

    missing_csrf = client.post(
        f"/api/v1/staff/payout-requests/{payout['id']}/reviews",
        json={"idempotency_key": str(uuid.uuid4())},
    )
    assert missing_csrf.status_code == 403
    review = _staff_command(
        client,
        payout["id"],
        "reviews",
        {"idempotency_key": str(uuid.uuid4()), "comment": "Documents received"},
    )
    assert review.status_code == 200, review.text
    review_body = review.json()
    assert set(review_body) == COMMAND_RECEIPT_FIELDS
    assert review_body["sequence_number"] == 2
    assert review_body["action"] == "review_started"
    assert review_body["from_status"] == "requested"
    assert review_body["status"] == "under_review"
    assert review_body["actor_user_id"] == str(manager_id)
    assert review_body["details"] == {"comment_recorded": True}

    no_requisites = _staff_command(
        client,
        payout["id"],
        "approvals",
        {
            "idempotency_key": str(uuid.uuid4()),
            "requisites_verified": False,
            "self_employment_verified": True,
        },
    )
    assert no_requisites.status_code == 422

    no_tax_check = _staff_command(
        client,
        payout["id"],
        "approvals",
        {
            "idempotency_key": str(uuid.uuid4()),
            "requisites_verified": True,
            "self_employment_verified": False,
        },
    )
    assert no_tax_check.status_code == 409
    assert no_tax_check.json()["error"]["code"] == "SELF_EMPLOYED_STATUS_REQUIRED"

    _login(client, finance_email)
    forbidden = _staff_command(
        client,
        payout["id"],
        "approvals",
        {
            "idempotency_key": str(uuid.uuid4()),
            "requisites_verified": True,
            "self_employment_verified": True,
        },
    )
    assert forbidden.status_code == 403

    _login(client, manager_email)
    approved = _staff_command(
        client,
        payout["id"],
        "approvals",
        {
            "idempotency_key": str(uuid.uuid4()),
            "requisites_verified": True,
            "self_employment_verified": True,
            "comment": "Requisites and tax status verified",
        },
    )
    assert approved.status_code == 200, approved.text
    approved_body = approved.json()
    assert set(approved_body) == COMMAND_RECEIPT_FIELDS
    assert approved_body["sequence_number"] == 3
    assert approved_body["action"] == "approved"
    assert approved_body["from_status"] == "under_review"
    assert approved_body["status"] == "approved"
    assert approved_body["actor_user_id"] == str(manager_id)
    assert approved_body["details"]["requisites_verified"] is True
    assert approved_body["details"]["self_employment_verified"] is True
    assert approved_body["details"]["comment_recorded"] is True

    stored = client.get(f"/api/v1/staff/payout-requests/{payout['id']}")
    assert stored.status_code == 200, stored.text
    stored_body = stored.json()
    assert stored_body["approved_by_user_id"] == str(manager_id)
    approved_at = datetime.fromisoformat(stored_body["approved_at"])
    assert date.fromisoformat(stored_body["payment_due_date"]) == (
        approved_at.astimezone(MOSCOW).date() + timedelta(days=20)
    )
    assert approved_body["details"]["payment_due_date"] == stored_body["payment_due_date"]
    assert stored_body["requisites_verified_at"] is not None
    assert stored_body["self_employment_verified_at"] is not None
    assert _balance(client, blogger_id).reserved_kopecks == 90_000


def test_rejection_requires_reason_releases_reserve_and_is_idempotent(client):
    blogger_id, _, payout = _prepare_request(
        client,
        label="rejection",
        available_kopecks=72_500,
    )
    manager_id, manager_email = _seed_staff(client, Role.MANAGER, "rejection-manager")
    _login(client, manager_email)
    reviewed = _staff_command(
        client,
        payout["id"],
        "reviews",
        {"idempotency_key": str(uuid.uuid4())},
    )
    assert reviewed.status_code == 200, reviewed.text

    invalid = _staff_command(
        client,
        payout["id"],
        "rejections",
        {"idempotency_key": str(uuid.uuid4()), "reason": "no"},
    )
    assert invalid.status_code == 422

    key = uuid.uuid4()
    payload = {
        "idempotency_key": str(key),
        "reason": "The payout details could not be verified",
        "comment": "Creator must update the SBP recipient",
    }
    rejected = _staff_command(client, payout["id"], "rejections", payload)
    assert rejected.status_code == 200, rejected.text
    body = rejected.json()
    assert set(body) == COMMAND_RECEIPT_FIELDS
    assert body["action"] == "rejected"
    assert body["from_status"] == "under_review"
    assert body["status"] == "rejected"
    assert body["actor_user_id"] == str(manager_id)
    assert body["details"] == {
        "reason_recorded": True,
        "comment_recorded": True,
    }
    stored = client.get(f"/api/v1/staff/payout-requests/{payout['id']}")
    assert stored.status_code == 200, stored.text
    assert stored.json()["rejection_reason"] == payload["reason"]
    assert stored.json()["rejected_by_user_id"] == str(manager_id)

    balance = _balance(client, blogger_id)
    assert (balance.available_kopecks, balance.reserved_kopecks, balance.paid_kopecks) == (
        72_500,
        0,
        0,
    )
    entries = _ledger(client, blogger_id)
    assert {_enum_value(entry.operation_type) for entry in entries} == {
        "payout_reserved",
        "payout_released",
    }
    released_entry = next(
        entry
        for entry in entries
        if _enum_value(entry.operation_type) == "payout_released"
    )
    assert (
        released_entry.available_delta_kopecks,
        released_entry.reserved_delta_kopecks,
        released_entry.paid_delta_kopecks,
    ) == (72_500, -72_500, 0)

    replay = _staff_command(client, payout["id"], "rejections", payload)
    assert replay.status_code == 200, replay.text
    assert replay.json() == body
    assert len(_ledger(client, blogger_id)) == 2

    reused = _staff_command(
        client,
        payout["id"],
        "rejections",
        dict(payload, reason="A different rejection reason"),
    )
    assert reused.status_code == 409
    assert reused.json()["error"]["code"] == "PAYOUT_IDEMPOTENCY_CONFLICT"

    with client.app.state.test_session() as db:
        events = list(
            db.scalars(
                select(PayoutEvent)
                .where(PayoutEvent.payout_request_id == uuid.UUID(payout["id"]))
                .order_by(PayoutEvent.sequence_number)
            )
        )
        assert [event.sequence_number for event in events] == [1, 2, 3]
        assert [_enum_value(event.action) for event in events] == [
            "requested",
            "review_started",
            "rejected",
        ]
        assert [_enum_value(event.to_status) for event in events] == [
            "requested",
            "under_review",
            "rejected",
        ]
        assert db.scalar(
            select(func.count())
            .select_from(SecurityEvent)
            .where(SecurityEvent.object_id == uuid.UUID(payout["id"]))
        ) == 3


def test_create_and_review_replays_return_original_receipts_after_approval(client):
    blogger_id, blogger_email = _seed_creator(
        client,
        label="historical-replay",
        available_kopecks=64_000,
    )
    manager_id, manager_email = _seed_staff(
        client,
        Role.MANAGER,
        "historical-replay-manager",
    )
    _login(client, blogger_email)
    assert _put_details(client).status_code == 200

    create_key = uuid.uuid4()
    created = _create_request(client, idempotency_key=create_key)
    assert created.status_code == 201, created.text
    create_receipt = created.json()
    assert create_receipt["actor_user_id"] == str(blogger_id)
    payout_id = create_receipt["payout_request_id"]

    review_key = uuid.uuid4()
    review_payload = {
        "idempotency_key": str(review_key),
        "comment": "Verified for historical replay",
    }
    _login(client, manager_email)
    reviewed = _staff_command(client, payout_id, "reviews", review_payload)
    assert reviewed.status_code == 200, reviewed.text
    review_receipt = reviewed.json()
    assert review_receipt["actor_user_id"] == str(manager_id)

    approved = _staff_command(
        client,
        payout_id,
        "approvals",
        {
            "idempotency_key": str(uuid.uuid4()),
            "requisites_verified": True,
            "self_employment_verified": None,
        },
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "approved"

    _login(client, blogger_email)
    create_replay = _create_request(client, idempotency_key=create_key)
    assert create_replay.status_code == 201, create_replay.text
    assert create_replay.json() == create_receipt
    assert create_replay.json()["status"] == "requested"

    _login(client, manager_email)
    review_replay = _staff_command(client, payout_id, "reviews", review_payload)
    assert review_replay.status_code == 200, review_replay.text
    assert review_replay.json() == review_receipt
    assert review_replay.json()["status"] == "under_review"


@pytest.mark.parametrize("role", [Role.MANAGER, Role.ADMIN])
def test_manager_or_admin_can_reject_requested_and_release_wallet(client, role):
    blogger_id, _, payout = _prepare_request(
        client,
        label=f"direct-rejection-{role.value}",
        available_kopecks=37_100,
    )
    staff_id, staff_email = _seed_staff(
        client,
        role,
        f"direct-rejection-{role.value}",
    )
    _login(client, staff_email)
    rejected = _staff_command(
        client,
        payout["id"],
        "rejections",
        {
            "idempotency_key": str(uuid.uuid4()),
            "reason": "The requested payout cannot be processed",
        },
    )
    assert rejected.status_code == 200, rejected.text
    receipt = rejected.json()
    assert receipt["sequence_number"] == 2
    assert receipt["action"] == "rejected"
    assert receipt["from_status"] == "requested"
    assert receipt["status"] == "rejected"
    assert receipt["actor_user_id"] == str(staff_id)

    balance = _balance(client, blogger_id)
    assert (balance.available_kopecks, balance.reserved_kopecks, balance.paid_kopecks) == (
        37_100,
        0,
        0,
    )
    assert {_enum_value(entry.operation_type) for entry in _ledger(client, blogger_id)} == {
        "payout_reserved",
        "payout_released",
    }


def test_payment_settles_reserved_money_and_finance_cannot_skip_approval(client):
    blogger_id, blogger_email, payout = _prepare_request(
        client,
        label="payment",
        available_kopecks=48_750,
    )
    _, manager_email = _seed_staff(client, Role.MANAGER, "payment-manager")
    finance_id, finance_email = _seed_staff(client, Role.FINANCE, "payment-finance")

    _login(client, finance_email)
    premature = _staff_command(
        client,
        payout["id"],
        "payments",
        {
            "idempotency_key": str(uuid.uuid4()),
            "paid_on": moscow_today().isoformat(),
            "payment_reference": "bank-early",
        },
    )
    assert premature.status_code == 409

    _review_and_approve(
        client,
        payout_id=payout["id"],
        manager_email=manager_email,
        self_employed=False,
    )
    forbidden = _staff_command(
        client,
        payout["id"],
        "payments",
        {
            "idempotency_key": str(uuid.uuid4()),
            "paid_on": moscow_today().isoformat(),
        },
    )
    assert forbidden.status_code == 403

    _login(client, finance_email)
    key = uuid.uuid4()
    payment_payload = {
        "idempotency_key": str(key),
        "paid_on": moscow_today().isoformat(),
        "payment_reference": "bank-order-42",
    }
    paid = _staff_command(client, payout["id"], "payments", payment_payload)
    assert paid.status_code == 200, paid.text
    body = paid.json()
    assert set(body) == COMMAND_RECEIPT_FIELDS
    assert body["action"] == "paid"
    assert body["from_status"] == "approved"
    assert body["status"] == "paid"
    assert body["actor_user_id"] == str(finance_id)
    assert body["details"] == {
        "paid_on": moscow_today().isoformat(),
        "payment_reference_recorded": True,
        "receipt_due_date": None,
    }

    stored = client.get(f"/api/v1/staff/payout-requests/{payout['id']}")
    assert stored.status_code == 200, stored.text
    stored_body = stored.json()
    assert stored_body["paid_on"] == moscow_today().isoformat()
    assert stored_body["paid_recorded_at"] is not None
    assert stored_body["paid_by_user_id"] == str(finance_id)
    assert stored_body["payment_reference"] == "bank-order-42"
    assert stored_body["receipt_status"] == "not_applicable"
    assert stored_body["receipt_due_date"] is None

    _login(client, blogger_email)
    my_detail = client.get(f"/api/v1/me/payout-requests/{payout['id']}")
    my_list = client.get("/api/v1/me/payout-requests")
    assert my_detail.status_code == 200, my_detail.text
    assert my_list.status_code == 200, my_list.text
    assert not (MY_PAYOUT_FORBIDDEN_FIELDS & my_detail.json().keys())
    assert not (MY_PAYOUT_FORBIDDEN_FIELDS & my_list.json()["items"][0].keys())
    assert my_detail.json()["status"] == "paid"
    assert my_detail.json()["paid_on"] == moscow_today().isoformat()

    balance = _balance(client, blogger_id)
    assert (balance.available_kopecks, balance.reserved_kopecks, balance.paid_kopecks) == (
        0,
        0,
        48_750,
    )
    entries = _ledger(client, blogger_id)
    assert {_enum_value(entry.operation_type) for entry in entries} == {
        "payout_reserved",
        "payout_paid",
    }
    paid_entry = next(
        entry
        for entry in entries
        if _enum_value(entry.operation_type) == "payout_paid"
    )
    assert (
        paid_entry.available_delta_kopecks,
        paid_entry.reserved_delta_kopecks,
        paid_entry.paid_delta_kopecks,
    ) == (0, -48_750, 48_750)

    _login(client, finance_email)
    replay = _staff_command(client, payout["id"], "payments", payment_payload)
    assert replay.status_code == 200, replay.text
    assert replay.json() == body
    assert len(_ledger(client, blogger_id)) == 2
    reused = _staff_command(
        client,
        payout["id"],
        "payments",
        dict(payment_payload, payment_reference="different-bank-order"),
    )
    assert reused.status_code == 409
    assert reused.json()["error"]["code"] == "PAYOUT_IDEMPOTENCY_CONFLICT"


def test_only_finance_can_reject_an_approved_unpaid_request(client):
    blogger_id, _, payout = _prepare_request(
        client,
        label="approved-rejection",
        available_kopecks=22_000,
    )
    _, manager_email = _seed_staff(client, Role.MANAGER, "approved-rejection-manager")
    finance_id, finance_email = _seed_staff(
        client,
        Role.FINANCE,
        "approved-rejection-finance",
    )
    _review_and_approve(
        client,
        payout_id=payout["id"],
        manager_email=manager_email,
        self_employed=False,
    )

    forbidden = _staff_command(
        client,
        payout["id"],
        "rejections",
        {
            "idempotency_key": str(uuid.uuid4()),
            "reason": "Manager must not cancel an approved request",
        },
    )
    assert forbidden.status_code == 403

    _login(client, finance_email)
    rejected = _staff_command(
        client,
        payout["id"],
        "rejections",
        {
            "idempotency_key": str(uuid.uuid4()),
            "reason": "The bank transfer was cancelled before execution",
        },
    )
    assert rejected.status_code == 200, rejected.text
    rejected_body = rejected.json()
    assert rejected_body["action"] == "rejected"
    assert rejected_body["from_status"] == "approved"
    assert rejected_body["status"] == "rejected"
    assert rejected_body["actor_user_id"] == str(finance_id)
    balance = _balance(client, blogger_id)
    assert (balance.available_kopecks, balance.reserved_kopecks, balance.paid_kopecks) == (
        22_000,
        0,
        0,
    )


def test_self_employed_receipt_deadline_blocks_only_while_overdue(
    client,
    monkeypatch,
):
    past_now = datetime(2026, 7, 1, 9, 0, tzinfo=timezone.utc)
    monkeypatch.setattr(payout_service, "utc_now", lambda: past_now)
    monkeypatch.setattr(payout_service, "moscow_today", lambda: date(2026, 7, 1))
    monkeypatch.setattr(payout_policy, "moscow_today", lambda: date(2026, 7, 1))
    blogger_id, blogger_email, payout = _prepare_request(
        client,
        label="receipt",
        available_kopecks=31_000,
        recipient_status=RecipientStatus.SELF_EMPLOYED,
    )
    _, manager_email = _seed_staff(client, Role.MANAGER, "receipt-manager")
    finance_id, finance_email = _seed_staff(client, Role.FINANCE, "receipt-finance")
    _review_and_approve(
        client,
        payout_id=payout["id"],
        manager_email=manager_email,
        self_employed=True,
    )

    paid_on = date(2026, 7, 1)
    _login(client, finance_email)
    paid = _staff_command(
        client,
        payout["id"],
        "payments",
        {
            "idempotency_key": str(uuid.uuid4()),
            "paid_on": paid_on.isoformat(),
            "payment_reference": "receipt-test-payment",
        },
    )
    assert paid.status_code == 200, paid.text
    assert paid.json()["details"]["receipt_due_date"] == (
        paid_on + timedelta(days=3)
    ).isoformat()
    paid_detail = client.get(f"/api/v1/staff/payout-requests/{payout['id']}")
    assert paid_detail.status_code == 200, paid_detail.text
    assert paid_detail.json()["receipt_status"] == "awaiting_receipt"

    overdue_today = date(2026, 7, 5)
    monkeypatch.setattr(payout_service, "moscow_today", lambda: overdue_today)
    monkeypatch.setattr(payout_policy, "moscow_today", lambda: overdue_today)
    overdue_detail = client.get(f"/api/v1/staff/payout-requests/{payout['id']}")
    assert overdue_detail.status_code == 200, overdue_detail.text
    assert overdue_detail.json()["receipt_status"] == "overdue"
    for params in (
        {"receiptStatus": "overdue"},
        {"receiptDueBefore": (paid_on + timedelta(days=3)).isoformat()},
        {"isReceiptOverdue": "true"},
    ):
        queue = client.get("/api/v1/staff/payout-requests", params=params)
        assert queue.status_code == 200, queue.text
        assert [item["id"] for item in queue.json()["items"]] == [payout["id"]]

    with client.app.state.test_session() as db:
        balance = db.get(CreatorBalance, blogger_id)
        balance.available_kopecks += 5_000
        db.add(
            BalanceLedgerEntry(
                blogger_id=blogger_id,
                operation_type="period_correction",
                available_delta_kopecks=5_000,
                reserved_delta_kopecks=0,
                paid_delta_kopecks=0,
                reference_type="test_earning",
                reference_id=uuid.uuid4(),
                idempotency_key=f"test-new-earning:{uuid.uuid4()}",
            )
        )
        db.commit()

    _login(client, blogger_email)
    blocked = _create_request(client)
    assert blocked.status_code == 409
    assert blocked.json()["error"]["code"] == "PAYOUT_BLOCKED_BY_OVERDUE_RECEIPT"
    assert _balance(client, blogger_id).available_kopecks == 5_000

    _login(client, finance_email)
    receipt_key = uuid.uuid4()
    receipt_payload = {
        "idempotency_key": str(receipt_key),
        "received_on": paid_on.isoformat(),
    }
    received = _staff_command(client, payout["id"], "receipts", receipt_payload)
    assert received.status_code == 200, received.text
    received_body = received.json()
    assert received_body["action"] == "receipt_received"
    assert received_body["from_status"] == "paid"
    assert received_body["status"] == "paid"
    assert received_body["actor_user_id"] == str(finance_id)
    assert received_body["details"] == {"received_on": paid_on.isoformat()}
    received_detail = client.get(f"/api/v1/staff/payout-requests/{payout['id']}")
    assert received_detail.status_code == 200, received_detail.text
    assert received_detail.json()["receipt_status"] == "received"
    assert received_detail.json()["receipt_received_on"] == paid_on.isoformat()
    assert received_detail.json()["receipt_recorded_at"] is not None
    assert received_detail.json()["receipt_received_by_user_id"] == str(finance_id)
    received_queue = client.get(
        "/api/v1/staff/payout-requests",
        params={"receiptStatus": "received"},
    )
    assert received_queue.status_code == 200, received_queue.text
    assert [item["id"] for item in received_queue.json()["items"]] == [payout["id"]]

    replay = _staff_command(client, payout["id"], "receipts", receipt_payload)
    assert replay.status_code == 200, replay.text
    assert replay.json() == received_body
    _login(client, blogger_email)
    next_request = _create_request(client)
    assert next_request.status_code == 201, next_request.text
    assert next_request.json()["details"] == {"amount_kopecks": 5_000}


def test_negative_correction_blocks_approval_and_payment_without_losing_reserve(client):
    blogger_id, _, payout = _prepare_request(
        client,
        label="negative-correction",
        available_kopecks=60_000,
    )
    _, manager_email = _seed_staff(client, Role.MANAGER, "negative-manager")
    _, finance_email = _seed_staff(client, Role.FINANCE, "negative-finance")

    with client.app.state.test_session() as db:
        balance = db.get(CreatorBalance, blogger_id)
        balance.available_kopecks = -100
        db.add(
            BalanceLedgerEntry(
                blogger_id=blogger_id,
                operation_type="period_correction",
                available_delta_kopecks=-100,
                reserved_delta_kopecks=0,
                paid_delta_kopecks=0,
                reference_type="accrual_correction",
                reference_id=uuid.uuid4(),
                idempotency_key=f"test-negative-approval:{uuid.uuid4()}",
            )
        )
        db.commit()

    _login(client, manager_email)
    reviewed = _staff_command(
        client,
        payout["id"],
        "reviews",
        {"idempotency_key": str(uuid.uuid4())},
    )
    assert reviewed.status_code == 200, reviewed.text
    blocked_approval = _staff_command(
        client,
        payout["id"],
        "approvals",
        {
            "idempotency_key": str(uuid.uuid4()),
            "requisites_verified": True,
            "self_employment_verified": None,
        },
    )
    assert blocked_approval.status_code == 409
    assert blocked_approval.json()["error"]["code"] == "PAYOUT_BLOCKED_BY_NEGATIVE_BALANCE"
    assert blocked_approval.json()["error"]["details"] == {}
    balance = _balance(client, blogger_id)
    assert (balance.available_kopecks, balance.reserved_kopecks) == (-100, 60_000)

    with client.app.state.test_session() as db:
        balance = db.get(CreatorBalance, blogger_id)
        balance.available_kopecks = 0
        db.add(
            BalanceLedgerEntry(
                blogger_id=blogger_id,
                operation_type="period_correction",
                available_delta_kopecks=100,
                reserved_delta_kopecks=0,
                paid_delta_kopecks=0,
                reference_type="accrual_correction",
                reference_id=uuid.uuid4(),
                idempotency_key=f"test-clear-negative:{uuid.uuid4()}",
            )
        )
        db.commit()

    approved = _staff_command(
        client,
        payout["id"],
        "approvals",
        {
            "idempotency_key": str(uuid.uuid4()),
            "requisites_verified": True,
            "self_employment_verified": None,
        },
    )
    assert approved.status_code == 200, approved.text

    with client.app.state.test_session() as db:
        balance = db.get(CreatorBalance, blogger_id)
        balance.available_kopecks = -50
        db.add(
            BalanceLedgerEntry(
                blogger_id=blogger_id,
                operation_type="period_correction",
                available_delta_kopecks=-50,
                reserved_delta_kopecks=0,
                paid_delta_kopecks=0,
                reference_type="accrual_correction",
                reference_id=uuid.uuid4(),
                idempotency_key=f"test-negative-payment:{uuid.uuid4()}",
            )
        )
        db.commit()

    _login(client, finance_email)
    blocked_payment = _staff_command(
        client,
        payout["id"],
        "payments",
        {
            "idempotency_key": str(uuid.uuid4()),
            "paid_on": moscow_today().isoformat(),
            "payment_reference": "must-not-be-recorded",
        },
    )
    assert blocked_payment.status_code == 409
    assert blocked_payment.json()["error"]["code"] == "PAYOUT_BLOCKED_BY_NEGATIVE_BALANCE"
    balance = _balance(client, blogger_id)
    assert (balance.available_kopecks, balance.reserved_kopecks, balance.paid_kopecks) == (
        -50,
        60_000,
        0,
    )
    with client.app.state.test_session() as db:
        stored = db.get(PayoutRequest, uuid.UUID(payout["id"]))
        assert _enum_value(stored.status) == "approved"
        assert stored.paid_on is None


def test_finance_sees_only_approved_and_paid_while_manager_sees_all(client):
    _, _, requested = _prepare_request(
        client,
        label="finance-queue-requested",
        available_kopecks=10_000,
    )
    _, _, approved = _prepare_request(
        client,
        label="finance-queue-approved",
        available_kopecks=20_000,
    )
    _, _, paid = _prepare_request(
        client,
        label="finance-queue-paid",
        available_kopecks=30_000,
    )
    _, manager_email = _seed_staff(client, Role.MANAGER, "finance-queue-manager")
    _, finance_email = _seed_staff(client, Role.FINANCE, "finance-queue-finance")

    _review_and_approve(
        client,
        payout_id=approved["id"],
        manager_email=manager_email,
        self_employed=False,
    )
    _review_and_approve(
        client,
        payout_id=paid["id"],
        manager_email=manager_email,
        self_employed=False,
    )
    _login(client, finance_email)
    payment = _staff_command(
        client,
        paid["id"],
        "payments",
        {
            "idempotency_key": str(uuid.uuid4()),
            "paid_on": moscow_today().isoformat(),
            "payment_reference": "finance-queue-payment",
        },
    )
    assert payment.status_code == 200, payment.text

    finance_list = client.get("/api/v1/staff/payout-requests?pageSize=100")
    assert finance_list.status_code == 200, finance_list.text
    _assert_private_no_store(finance_list)
    assert finance_list.json()["total_items"] == 2
    assert {item["status"] for item in finance_list.json()["items"]} == {
        "approved",
        "paid",
    }
    requested_queue = client.get(
        "/api/v1/staff/payout-requests?status=requested&pageSize=100"
    )
    assert requested_queue.status_code == 200, requested_queue.text
    assert requested_queue.json()["items"] == []
    assert requested_queue.json()["total_items"] == 0
    assert (
        client.get(f"/api/v1/staff/payout-requests/{requested['id']}").status_code
        == 404
    )
    assert client.get(f"/api/v1/staff/payout-requests/{approved['id']}").status_code == 200
    assert client.get(f"/api/v1/staff/payout-requests/{paid['id']}").status_code == 200

    by_number = client.get(
        "/api/v1/staff/payout-requests",
        params={"requestNumber": approved["request_number"]},
    )
    assert by_number.status_code == 200, by_number.text
    assert [item["id"] for item in by_number.json()["items"]] == [approved["id"]]

    _login(client, manager_email)
    manager_list = client.get("/api/v1/staff/payout-requests?pageSize=100")
    assert manager_list.status_code == 200, manager_list.text
    assert manager_list.json()["total_items"] == 3
    assert {item["status"] for item in manager_list.json()["items"]} == {
        "requested",
        "approved",
        "paid",
    }
    requested_detail = client.get(f"/api/v1/staff/payout-requests/{requested['id']}")
    assert requested_detail.status_code == 200, requested_detail.text
    _assert_private_no_store(requested_detail)

    approved_on = datetime.fromisoformat(
        client.get(f"/api/v1/staff/payout-requests/{approved['id']}").json()[
            "approved_at"
        ]
    ).astimezone(MOSCOW).date().isoformat()
    approved_range = client.get(
        "/api/v1/staff/payout-requests",
        params={
            "approvedFrom": approved_on,
            "approvedTo": approved_on,
            "pageSize": 100,
        },
    )
    assert approved_range.status_code == 200, approved_range.text
    assert {item["id"] for item in approved_range.json()["items"]} == {
        approved["id"],
        paid["id"],
    }
    invalid_range = client.get(
        "/api/v1/staff/payout-requests",
        params={"approvedFrom": "2026-08-22", "approvedTo": "2026-08-21"},
    )
    assert invalid_range.status_code == 422


def test_owner_scope_and_pagination_apply_to_creator_and_staff_lists(client):
    first_id, first_email, first = _prepare_request(
        client,
        label="owner-first",
        available_kopecks=1_000,
    )
    second_id, second_email, second = _prepare_request(
        client,
        label="owner-second",
        available_kopecks=2_000,
        recipient_status=RecipientStatus.SELF_EMPLOYED,
    )

    _login(client, first_email)
    mine = client.get("/api/v1/me/payout-requests?page=1&pageSize=1&status=requested")
    assert mine.status_code == 200, mine.text
    assert mine.json()["total_items"] == 1
    assert mine.json()["total_pages"] == 1
    assert mine.json()["page_size"] == 1
    assert mine.json()["items"][0]["id"] == first["id"]
    assert client.get(f"/api/v1/me/payout-requests/{second['id']}").status_code == 404

    _login(client, second_email)
    assert client.get(f"/api/v1/me/payout-requests/{first['id']}").status_code == 404
    assert client.get(f"/api/v1/me/payout-requests/{second['id']}").status_code == 200

    _, manager_email = _seed_staff(client, Role.MANAGER, "list-manager")
    _login(client, manager_email)
    first_page = client.get("/api/v1/staff/payout-requests?page=1&pageSize=1&status=requested")
    second_page = client.get("/api/v1/staff/payout-requests?page=2&pageSize=1&status=requested")
    assert first_page.status_code == 200, first_page.text
    assert second_page.status_code == 200, second_page.text
    assert first_page.json()["total_items"] == 2
    assert first_page.json()["total_pages"] == 2
    returned_ids = {
        first_page.json()["items"][0]["blogger_id"],
        second_page.json()["items"][0]["blogger_id"],
    }
    assert returned_ids == {str(first_id), str(second_id)}

    filtered = client.get(
        "/api/v1/staff/payout-requests"
        f"?recipientType=self_employed&bloggerId={second_id}&pageSize=10"
    )
    assert filtered.status_code == 200, filtered.text
    assert filtered.json()["total_items"] == 1
    assert filtered.json()["items"][0]["id"] == second["id"]

    detail = client.get(f"/api/v1/staff/payout-requests/{first['id']}")
    assert detail.status_code == 200, detail.text
    assert [event["action"] for event in detail.json()["history"]] == ["requested"]
    assert "idempotency_key" not in detail.json()["history"][0]
    assert "payload_hash" not in detail.json()["history"][0]


def _assert_not_a_spreadsheet_formula(value: str) -> None:
    assert value.lstrip()[:1] not in {"=", "+", "-", "@"}


def _xlsx_strings_and_formula_count(content: bytes) -> tuple[list[str], int]:
    strings: list[str] = []
    formula_count = 0
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for name in archive.namelist():
            if not name.startswith("xl/") or not name.endswith(".xml"):
                continue
            root = ElementTree.fromstring(archive.read(name))
            for element in root.iter():
                local_name = element.tag.rsplit("}", 1)[-1]
                if local_name == "f":
                    formula_count += 1
                elif local_name == "t" and element.text is not None:
                    strings.append(element.text)
    return strings, formula_count


class _MemoryArtifactStore:
    def __init__(self):
        self.objects = {}

    def upload(self, *, key, path, content_type, content_sha256):
        self.objects[key] = (path.read_bytes(), content_type, content_sha256)

    def download(self, *, key):
        content, content_type, _ = self.objects[key]
        return ArtifactDownload(io.BytesIO(content), len(content), content_type)

    def delete(self, *, key):
        self.objects.pop(key, None)


class _ExportProducer:
    def __init__(self):
        self.commands = []

    def send_task(self, task_name, *, args, queue):
        self.commands.append((task_name, args[0], queue))


def _complete_export(client, *, export_format, store, idempotency_key=None):
    today = moscow_today()
    created = client.post(
        "/api/v1/staff/payout-exports",
        headers=_csrf_headers(client),
        json={
            "idempotency_key": str(idempotency_key or uuid.uuid4()),
            "export_type": "payout_register",
            "format": export_format,
            "filters": {
                "requested_from": today.replace(day=1).isoformat(),
                "requested_to": today.isoformat(),
            },
        },
    )
    assert created.status_code == 202, created.text
    producer = _ExportProducer()
    assert dispatch_export(
        session_factory=client.app.state.test_session,
        task_producer=producer,
    )
    command = ExportCommand.model_validate(producer.commands[0][1])
    execute_export(command, client.app.state.test_session, store)
    status_response = client.get(
        f"/api/v1/staff/payout-exports/{created.json()['id']}"
    )
    return created, status_response


def _export_payload(*, export_type="payout_register", export_format="csv", key=None):
    today = moscow_today()
    return {
        "idempotency_key": str(key or uuid.uuid4()),
        "export_type": export_type,
        "format": export_format,
        "filters": {
            "requested_from": today.replace(day=1).isoformat(),
            "requested_to": today.isoformat(),
        },
    }


def test_export_api_enforces_roles_csrf_ranges_and_idempotency(client):
    _, manager_email = _seed_staff(client, Role.MANAGER, "export-denied-manager")
    _login(client, manager_email)
    denied = client.post(
        "/api/v1/staff/payout-exports",
        headers=_csrf_headers(client),
        json=_export_payload(),
    )
    assert denied.status_code == 403

    _, finance_email = _seed_staff(client, Role.FINANCE, "export-contract-finance")
    _login(client, finance_email)
    missing_csrf = client.post(
        "/api/v1/staff/payout-exports", json=_export_payload()
    )
    assert missing_csrf.status_code == 403
    history_denied = client.post(
        "/api/v1/staff/payout-exports",
        headers=_csrf_headers(client),
        json=_export_payload(export_type="payout_history"),
    )
    assert history_denied.status_code == 403

    idempotency_key = uuid.uuid4()
    payload = _export_payload(key=idempotency_key)
    first = client.post(
        "/api/v1/staff/payout-exports",
        headers=_csrf_headers(client),
        json=payload,
    )
    repeated = client.post(
        "/api/v1/staff/payout-exports",
        headers=_csrf_headers(client),
        json=payload,
    )
    assert first.status_code == repeated.status_code == 202
    assert first.json()["id"] == repeated.json()["id"]
    conflicting = client.post(
        "/api/v1/staff/payout-exports",
        headers=_csrf_headers(client),
        json={**payload, "format": "xlsx"},
    )
    assert conflicting.status_code == 409
    assert conflicting.json()["error"]["code"] == "EXPORT_IDEMPOTENCY_CONFLICT"

    too_wide = _export_payload()
    too_wide["filters"] = {
        "requested_from": "2025-01-01",
        "requested_to": "2026-02-01",
    }
    invalid_range = client.post(
        "/api/v1/staff/payout-exports",
        headers=_csrf_headers(client),
        json=too_wide,
    )
    assert invalid_range.status_code == 422

    _, admin_email = _seed_staff(client, Role.ADMIN, "export-history-admin")
    _login(client, admin_email)
    history = client.post(
        "/api/v1/staff/payout-exports",
        headers=_csrf_headers(client),
        json=_export_payload(export_type="payout_history"),
    )
    assert history.status_code == 202
    assert history.json()["export_type"] == "payout_history"


def test_export_dispatch_broker_failure_releases_lease(client):
    _, finance_email = _seed_staff(client, Role.FINANCE, "export-broker-finance")
    _login(client, finance_email)
    created = client.post(
        "/api/v1/staff/payout-exports",
        headers=_csrf_headers(client),
        json=_export_payload(),
    )
    assert created.status_code == 202

    class FailingProducer:
        def send_task(self, *args, **kwargs):
            del args, kwargs
            raise RuntimeError("broker unavailable")

    assert not dispatch_export(
        session_factory=client.app.state.test_session,
        task_producer=FailingProducer(),
    )
    status_response = client.get(
        f"/api/v1/staff/payout-exports/{created.json()['id']}"
    )
    assert status_response.status_code == 200
    assert status_response.json()["status"] == "retry_wait"
    assert status_response.json()["attempt_count"] == 0
    assert status_response.json()["error_code"] == "broker_publish_failed"


def test_csv_and_xlsx_exports_are_filtered_auditable_and_formula_safe(client):
    dangerous_name = "  =HYPERLINK(\"https://example.test\")"
    dangerous_bank = "@SUM(1+1)"
    _, _, payout = _prepare_request(
        client,
        label="export",
        available_kopecks=12_345,
        full_name=dangerous_name,
        display_name="+Dangerous display",
        bank_name=dangerous_bank,
    )
    _, manager_email = _seed_staff(client, Role.MANAGER, "export-manager")
    _review_and_approve(
        client,
        payout_id=payout["id"],
        manager_email=manager_email,
        self_employed=False,
    )
    _, _, one_ruble_payout = _prepare_request(
        client,
        label="export-one-ruble",
        available_kopecks=100,
    )
    _review_and_approve(
        client,
        payout_id=one_ruble_payout["id"],
        manager_email=manager_email,
        self_employed=False,
    )
    _, finance_email = _seed_staff(client, Role.FINANCE, "export-finance")
    _login(client, finance_email)
    store = _MemoryArtifactStore()
    client.app.dependency_overrides[get_artifact_store] = lambda: store

    _, csv_status = _complete_export(client, export_format="csv", store=store)
    assert csv_status.status_code == 200, csv_status.text
    assert csv_status.json()["status"] == "ready"
    csv_response = client.get(csv_status.json()["download_url"])
    assert csv_response.status_code == 200, csv_response.text
    _assert_private_no_store(csv_response)
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert re.fullmatch(
        'attachment; filename="payout-register-[0-9]{4}-[0-9]{2}-[0-9]{2}-[a-f0-9]{8}.csv"',
        csv_response.headers["content-disposition"],
    )
    assert csv_response.headers["x-content-type-options"] == "nosniff"
    rows = list(
        csv.DictReader(
            io.StringIO(csv_response.content.decode("utf-8-sig")), delimiter=";"
        )
    )
    assert len(rows) == 2
    rows_by_number = {row["Номер заявки"]: row for row in rows}
    dangerous_row = rows_by_number[payout["request_number"]]
    one_ruble_row = rows_by_number[one_ruble_payout["request_number"]]
    assert dangerous_row["Сумма, руб."] == "123,45"
    assert one_ruble_row["Сумма, руб."] == "1,00"
    for row in rows:
        for column in ("Дата заявки", "Дата одобрения", "Срок оплаты"):
            assert re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", row[column])
    assert "HYPERLINK" in dangerous_row["Получатель"]
    assert "SUM" in dangerous_row["Банк"]
    _assert_not_a_spreadsheet_formula(dangerous_row["Получатель"])
    _assert_not_a_spreadsheet_formula(dangerous_row["Банк"])

    _, xlsx_status = _complete_export(client, export_format="xlsx", store=store)
    assert xlsx_status.status_code == 200, xlsx_status.text
    assert xlsx_status.json()["status"] == "ready"
    xlsx_response = client.get(xlsx_status.json()["download_url"])
    assert xlsx_response.status_code == 200, xlsx_response.text
    _assert_private_no_store(xlsx_response)
    assert xlsx_response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert re.fullmatch(
        'attachment; filename="payout-register-[0-9]{4}-[0-9]{2}-[0-9]{2}-[a-f0-9]{8}.xlsx"',
        xlsx_response.headers["content-disposition"],
    )
    strings, formula_count = _xlsx_strings_and_formula_count(xlsx_response.content)
    assert formula_count == 0
    exported_name = next(value for value in strings if "HYPERLINK" in value)
    exported_bank = next(value for value in strings if "SUM" in value)
    _assert_not_a_spreadsheet_formula(exported_name)
    _assert_not_a_spreadsheet_formula(exported_bank)

    workbook = load_workbook(
        io.BytesIO(xlsx_response.content),
        read_only=True,
        data_only=False,
    )
    worksheet = workbook.active
    xlsx_rows = list(worksheet.iter_rows())
    columns = {cell.value: index for index, cell in enumerate(xlsx_rows[0])}
    data_rows = {row[columns["Номер заявки"]].value: row for row in xlsx_rows[1:]}
    one_ruble_cells = data_rows[one_ruble_payout["request_number"]]
    amount_cell = one_ruble_cells[columns["Сумма, руб."]]
    assert amount_cell.value == 1
    assert amount_cell.number_format == "0.00"
    for row in data_rows.values():
        for column in ("Дата заявки", "Дата одобрения", "Срок оплаты"):
            cell = row[columns[column]]
            assert isinstance(cell.value, (date, datetime))
            assert cell.number_format.lower() == "dd.mm.yyyy"
            assert cell.data_type != "f"
    workbook.close()

    with client.app.state.test_session() as db:
        audit_events = list(
            db.scalars(
                select(SecurityEvent)
                .where(
                    SecurityEvent.action == AuditAction.PAYOUT_EXPORT_DOWNLOADED,
                    SecurityEvent.actor_user_id.is_not(None),
                )
                .order_by(SecurityEvent.occurred_at, SecurityEvent.id)
            )
        )
        assert len(audit_events) == 2
        assert {event.event_metadata["format"] for event in audit_events} == {
            "csv",
            "xlsx",
        }
        events_by_format = {
            event.event_metadata["format"]: event for event in audit_events
        }
        response_by_format = {"csv": csv_response, "xlsx": xlsx_response}
        for export_format, event in events_by_format.items():
            metadata = event.event_metadata
            assert metadata["row_count"] == 2
            assert metadata["content_sha256"] == hashlib.sha256(
                response_by_format[export_format].content
            ).hexdigest()
            assert event.object_type == "export_job"


def test_payout_snapshot_and_history_are_immutable(client):
    _, _, payout = _prepare_request(
        client,
        label="immutable",
        available_kopecks=8_000,
    )
    payout_id = uuid.UUID(payout["id"])

    with client.app.state.test_session() as db:
        stored = db.get(PayoutRequest, payout_id)
        stored.amount_kopecks = 1
        with pytest.raises(ValueError, match="snapshot is immutable"):
            db.flush()
        db.rollback()

    with client.app.state.test_session() as db:
        event = db.scalar(
            select(PayoutEvent).where(PayoutEvent.payout_request_id == payout_id)
        )
        event.event_metadata = {"rewritten": True}
        with pytest.raises(ValueError, match="append-only"):
            db.flush()
        db.rollback()

    with client.app.state.test_session() as db:
        event = db.scalar(
            select(PayoutEvent).where(PayoutEvent.payout_request_id == payout_id)
        )
        db.delete(event)
        with pytest.raises(ValueError, match="append-only"):
            db.flush()
        db.rollback()


def test_promoted_blogger_cannot_process_own_payout(client):
    blogger_id, blogger_email, payout = _prepare_request(
        client,
        label="self-processing",
        available_kopecks=32_100,
    )
    with client.app.state.test_session() as db:
        blogger = db.get(User, blogger_id)
        blogger.role = Role.ADMIN
        db.commit()

    _login(client, blogger_email)
    response = _staff_command(
        client,
        payout["id"],
        "reviews",
        {"idempotency_key": str(uuid.uuid4())},
    )

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "PAYOUT_SELF_PROCESSING_FORBIDDEN"
    balance = _balance(client, blogger_id)
    assert (balance.available_kopecks, balance.reserved_kopecks) == (0, 32_100)
    with client.app.state.test_session() as db:
        stored = db.get(PayoutRequest, uuid.UUID(payout["id"]))
        assert _enum_value(stored.status) == "requested"
        assert db.scalar(
            select(func.count()).select_from(PayoutEvent).where(
                PayoutEvent.payout_request_id == stored.id
            )
        ) == 1


def test_staff_history_preserves_each_command_comment(client):
    _, _, payout = _prepare_request(
        client,
        label="event-details",
        available_kopecks=43_200,
    )
    _, manager_email = _seed_staff(client, Role.MANAGER, "event-details-manager")
    _login(client, manager_email)
    reviewed = _staff_command(
        client,
        payout["id"],
        "reviews",
        {
            "idempotency_key": str(uuid.uuid4()),
            "comment": "Initial document review",
        },
    )
    assert reviewed.status_code == 200, reviewed.text
    approved = _staff_command(
        client,
        payout["id"],
        "approvals",
        {
            "idempotency_key": str(uuid.uuid4()),
            "requisites_verified": True,
            "comment": "Final approval note",
        },
    )
    assert approved.status_code == 200, approved.text

    detail = client.get(f"/api/v1/staff/payout-requests/{payout['id']}")
    assert detail.status_code == 200, detail.text
    comments = {
        event["action"]: event["comment"]
        for event in detail.json()["history"]
        if event["comment"] is not None
    }
    assert comments == {
        "review_started": "Initial document review",
        "approved": "Final approval note",
    }
    with client.app.state.test_session() as db:
        assert db.scalar(select(func.count()).select_from(PayoutEventDetail)) == 2

        stored = db.get(PayoutRequest, uuid.UUID(payout["id"]))
        stored.approved_at = stored.approved_at + timedelta(seconds=1)
        with pytest.raises(ValueError, match="fact is immutable"):
            db.flush()
        db.rollback()

    with client.app.state.test_session() as db:
        detail_row = db.scalar(
            select(PayoutEventDetail).where(
                PayoutEventDetail.comment == "Initial document review"
            )
        )
        detail_row.comment = "Rewritten review"
        with pytest.raises(ValueError, match="append-only"):
            db.flush()
        db.rollback()


def test_payout_retention_anonymizes_only_expired_terminal_data(client):
    blogger_id, _, payout = _prepare_request(
        client,
        label="retention",
        available_kopecks=65_400,
        full_name="Sensitive Creator",
        display_name="Sensitive Channel",
        bank_name="Sensitive Bank",
    )
    _, manager_email = _seed_staff(client, Role.MANAGER, "retention-manager")
    _login(client, manager_email)
    reviewed = _staff_command(
        client,
        payout["id"],
        "reviews",
        {
            "idempotency_key": str(uuid.uuid4()),
            "comment": "Contains private context",
        },
    )
    assert reviewed.status_code == 200, reviewed.text
    approved = _staff_command(
        client,
        payout["id"],
        "approvals",
        {
            "idempotency_key": str(uuid.uuid4()),
            "requisites_verified": True,
            "comment": "Approval private context",
        },
    )
    assert approved.status_code == 200, approved.text
    _, finance_email = _seed_staff(client, Role.FINANCE, "retention-finance")
    _login(client, finance_email)
    paid = _staff_command(
        client,
        payout["id"],
        "payments",
        {
            "idempotency_key": str(uuid.uuid4()),
            "paid_on": approved.json()["recorded_at"][:10],
            "payment_reference": "BANK-REFERENCE-KEPT",
        },
    )
    assert paid.status_code == 200, paid.text

    payout_id = uuid.UUID(payout["id"])
    old_time = datetime(2020, 1, 10, tzinfo=timezone.utc)
    anonymized_at = datetime(2026, 8, 23, tzinfo=timezone.utc)
    with client.app.state.test_session() as db:
        db.execute(
            update(User)
            .where(User.id == blogger_id)
            .values(
                status=AccountStatus.DELETED,
                status_changed_at=old_time,
                collaboration_ended_at=old_time,
            )
        )
        db.execute(
            update(PayoutRequest)
            .where(PayoutRequest.id == payout_id)
            .values(requested_at=old_time, paid_on=old_time.date())
        )
        db.commit()

    with client.app.state.test_session.begin() as db:
        result = anonymize_eligible_payout_pii(
            db,
            cutoff=date(2021, 1, 1),
            anonymized_at=anonymized_at,
            audit_context=AuditContext(
                request_id="retention-test",
                ip_address="127.0.0.1",
                user_agent="pytest",
            ),
        )
    assert result.bloggers == 1
    assert result.payout_requests == 1
    assert result.payout_details == 1
    assert result.event_details == 3

    with client.app.state.test_session() as db:
        stored = db.get(PayoutRequest, payout_id)
        details = db.get(PayoutDetails, blogger_id)
        event_details = list(
            db.scalars(
                select(PayoutEventDetail)
                .join(PayoutEvent)
                .where(PayoutEvent.payout_request_id == payout_id)
            )
        )
        audit = db.scalar(
            select(SecurityEvent).where(
                SecurityEvent.action == AuditAction.PAYOUT_PII_ANONYMIZED,
                SecurityEvent.object_id == blogger_id,
            )
        )
        assert stored.recipient_full_name == ANONYMIZED_TEXT
        assert stored.recipient_display_name == ANONYMIZED_TEXT
        assert stored.sbp_phone == ANONYMIZED_TEXT
        assert stored.bank_name is None
        assert stored.manager_comment is None
        assert stored.payment_reference == "BANK-REFERENCE-KEPT"
        assert stored.amount_kopecks == 65_400
        assert stored.pii_anonymized_at.replace(tzinfo=timezone.utc) == anonymized_at
        assert details.sbp_phone == ANONYMIZED_TEXT
        assert details.bank_name is None
        assert all(
            item.pii_anonymized_at.replace(tzinfo=timezone.utc) == anonymized_at
            for item in event_details
        )
        assert {item.comment for item in event_details} == {ANONYMIZED_TEXT, None}
        assert audit is not None

    with client.app.state.test_session.begin() as db:
        replay = anonymize_eligible_payout_pii(
            db,
            cutoff=date(2021, 1, 1),
            anonymized_at=anonymized_at,
            audit_context=AuditContext(
                request_id="retention-replay",
                ip_address="127.0.0.1",
                user_agent="pytest",
            ),
        )
    assert replay.bloggers == 0


def test_export_treats_naive_database_timestamps_as_utc():
    payout = SimpleNamespace(
        request_number="PAY-20260101-AAAAAAAAAAAAAAAAAAA",
        recipient_full_name="Boundary Creator",
        recipient_display_name="Boundary",
        recipient_type="individual",
        sbp_phone=VALID_SBP_PHONE,
        bank_name=None,
        amount_kopecks=100,
        requested_at=datetime(2026, 1, 1, 20, 59, 59),
        approved_at=datetime(2026, 1, 1, 21, 0, 0),
        payment_due_date=date(2026, 1, 21),
        self_employment_verified_at=None,
        manager_comment=None,
        status="approved",
    )

    assert moscow_date(payout.requested_at) == date(2026, 1, 1)
    assert moscow_date(payout.approved_at) == date(2026, 1, 2)


def test_staff_list_captures_moscow_today_once_for_filter_and_response(
    client,
    monkeypatch,
):
    _, _, payout = _prepare_request(
        client,
        label="single-clock",
        available_kopecks=54_300,
    )
    _, manager_email = _seed_staff(client, Role.MANAGER, "single-clock-manager")
    _review_and_approve(
        client,
        payout_id=payout["id"],
        manager_email=manager_email,
        self_employed=False,
    )
    effective_today = date(2026, 8, 23)
    with client.app.state.test_session() as db:
        db.execute(
            update(PayoutRequest)
            .where(PayoutRequest.id == uuid.UUID(payout["id"]))
            .values(payment_due_date=effective_today - timedelta(days=1))
        )
        db.commit()

    calls = 0

    def one_today():
        nonlocal calls
        calls += 1
        if calls > 1:
            raise AssertionError("moscow_today was called more than once")
        return effective_today

    monkeypatch.setattr(payout_service, "moscow_today", one_today)
    monkeypatch.setattr(
        payout_policy,
        "moscow_today",
        lambda: (_ for _ in ()).throw(AssertionError("DTO read the clock again")),
    )
    response = client.get("/api/v1/staff/payout-requests?isOverdue=true")
    assert response.status_code == 200, response.text
    assert calls == 1
    assert [item["id"] for item in response.json()["items"]] == [payout["id"]]
    assert response.json()["items"][0]["is_payment_overdue"] is True


def test_payout_register_rejects_selection_above_sync_limit(client, monkeypatch):
    _, _, payout = _prepare_request(
        client,
        label="export-limit",
        available_kopecks=65_400,
    )
    _, manager_email = _seed_staff(client, Role.MANAGER, "export-limit-manager")
    _review_and_approve(
        client,
        payout_id=payout["id"],
        manager_email=manager_email,
        self_employed=False,
    )
    _, finance_email = _seed_staff(client, Role.FINANCE, "export-limit-finance")
    _login(client, finance_email)
    monkeypatch.setattr(export_payout_data, "MAX_EXPORT_ROWS", 0)
    store = _MemoryArtifactStore()

    _, response = _complete_export(client, export_format="csv", store=store)

    assert response.status_code == 200
    assert response.json()["status"] == "failed"
    assert response.json()["error_code"] == "export_too_large"
    assert response.json()["download_url"] is None
    _assert_private_no_store(response)
