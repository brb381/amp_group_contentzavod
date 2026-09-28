import io
import uuid
import zipfile
from datetime import date, datetime, timedelta, timezone

import pytest
from sqlalchemy import select

from app.auth.models import AccountStatus, Role, User
from app.auth.security import hash_password
from app.clock import utc_now
from app.catalog.models import Brand
from app.contracts import ExportCommand
from app.content.models import (
    Publication,
    PublicationParseStatus,
    PublicationStatus,
    VideoCard,
)
from app.creators.models import CreatorProfile, SocialAccount, SocialAccountStatus
from app.exports.data import load_data_export_rows
from app.exports.models import ExportFormat, ExportJob, ExportType
from app.exports.generator import generate_export, schema_for
from app.exports.processor import cleanup_expired_artifacts, execute_export
from app.exports.schemas import DataExportFilters
from app.exports.storage import ArtifactDownload
from app.platforms import Platform
from app.readings.models import ReadingSource, ReadingStatus, ViewReading
from app.support.models import SupportCategory, SupportMessage, SupportStatus, SupportTicket
from app.scheduling.exports import dispatch_export


PASSWORD = "export-test-password-123"
DATA_FILTERS = {"date_from": "2026-01-01", "date_to": "2026-09-01"}
PAYOUT_FILTERS = {
    "requested_from": "2026-01-01",
    "requested_to": "2026-09-01",
}


class MemoryArtifactStore:
    def __init__(self):
        self.objects = {}

    def upload(self, *, key, path, content_type, content_sha256):
        self.objects[key] = (path.read_bytes(), content_type, content_sha256)

    def download(self, *, key):
        content, content_type, _ = self.objects[key]
        return ArtifactDownload(io.BytesIO(content), len(content), content_type)

    def delete(self, *, key):
        self.objects.pop(key, None)


class CapturingProducer:
    def __init__(self):
        self.command = None

    def send_task(self, task_name, *, args, queue):
        del task_name, queue
        self.command = ExportCommand.model_validate(args[0])


def _seed_staff(client, role: Role, label: str) -> str:
    email = f"export-{label}-{uuid.uuid4()}@example.com"
    with client.app.state.test_session() as db:
        db.add(
            User(
                email=email,
                password_hash=hash_password(PASSWORD),
                role=role,
                status=AccountStatus.ACTIVE,
                email_verified_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
            )
        )
        db.commit()
    return email


def _login(client, email: str) -> None:
    response = client.post(
        "/api/v1/auth/login", json={"email": email, "password": PASSWORD}
    )
    assert response.status_code == 200, response.text


def _payload(export_type: ExportType, *, export_format: str = "csv") -> dict:
    filters = dict(
        PAYOUT_FILTERS if export_type in {
            ExportType.PAYOUT_REGISTER, ExportType.PAYOUT_HISTORY
        } else DATA_FILTERS
    )
    return {
        "idempotency_key": str(uuid.uuid4()),
        "export_type": export_type.value,
        "format": export_format,
        "filters": filters,
    }


def _post_export(client, export_type: ExportType):
    return client.post(
        "/api/v1/staff/exports",
        headers={"X-CSRF-Token": client.cookies.get("amp_csrf")},
        json=_payload(export_type),
    )


@pytest.mark.parametrize(
    ("role", "export_type"),
    [
        (Role.ADMIN, ExportType.AUDIT_LOG),
        (Role.FINANCE, ExportType.PAYOUT_HISTORY),
        (Role.MANAGER, ExportType.BLOGGERS),
        (Role.MODERATOR, ExportType.MODERATION_HISTORY),
        (Role.ANALYST, ExportType.ACCRUALS),
    ],
)
def test_export_role_matrix_allows_expected_types(client, role, export_type):
    _login(client, _seed_staff(client, role, f"allowed-{role.value}"))

    response = _post_export(client, export_type)

    assert response.status_code == 202, response.text
    assert response.json()["export_type"] == export_type.value


@pytest.mark.parametrize(
    ("role", "export_type"),
    [
        (Role.FINANCE, ExportType.PUBLICATIONS),
        (Role.MANAGER, ExportType.PAYOUT_REGISTER),
        (Role.MODERATOR, ExportType.BLOGGERS),
        (Role.ANALYST, ExportType.BLOGGERS),
        (Role.ANALYST, ExportType.AUDIT_LOG),
    ],
)
def test_export_role_matrix_rejects_sensitive_types(client, role, export_type):
    _login(client, _seed_staff(client, role, f"denied-{role.value}"))

    response = _post_export(client, export_type)

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "EXPORT_FORBIDDEN"


def test_every_export_type_runs_through_worker_with_code_generated_csv(client):
    _login(client, _seed_staff(client, Role.ADMIN, "all-types"))
    store = MemoryArtifactStore()

    for export_type in ExportType:
        created = _post_export(client, export_type)
        assert created.status_code == 202, created.text
        producer = CapturingProducer()
        assert dispatch_export(
            session_factory=client.app.state.test_session,
            task_producer=producer,
        )
        execute_export(producer.command, client.app.state.test_session, store)

        with client.app.state.test_session() as db:
            export_id = uuid.UUID(created.json()["id"])
            job = db.scalar(select(ExportJob).where(ExportJob.id == export_id))
            assert job.status.value == "ready", (export_type, job.last_error_code)
            assert job.artifact_key.startswith("exports/")
            assert job.row_count >= 0
            content = store.objects[job.artifact_key][0]
            assert content.startswith(b"\xef\xbb\xbf")

    deleted = cleanup_expired_artifacts(
        client.app.state.test_session,
        store,
        now=utc_now() + timedelta(days=366),
    )
    assert deleted == len(ExportType)
    assert store.objects == {}
    with client.app.state.test_session() as db:
        jobs = list(db.scalars(select(ExportJob)))
        assert all(job.status.value == "expired" for job in jobs)
        assert all(job.artifact_deleted_at is not None for job in jobs)


def test_general_export_validates_type_specific_filters(client):
    _login(client, _seed_staff(client, Role.ADMIN, "filter-validation"))
    payload = _payload(ExportType.BLOGGERS)
    payload["filters"]["platform"] = "vk"

    response = client.post(
        "/api/v1/staff/exports",
        headers={"X-CSRF-Token": client.cookies.get("amp_csrf")},
        json=payload,
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "EXPORT_FILTER_NOT_SUPPORTED"


def test_social_account_export_includes_non_approved_statuses(client):
    with client.app.state.test_session() as db:
        blogger = User(
            email=f"social-export-{uuid.uuid4()}@example.com",
            password_hash=hash_password(PASSWORD),
            role=Role.BLOGGER,
            status=AccountStatus.ACTIVE,
            email_verified_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
        )
        db.add(blogger)
        db.flush()
        db.add(CreatorProfile(user_id=blogger.id, display_name="Тестовый блогер"))
        db.add_all([
            SocialAccount(
                user_id=blogger.id, platform=Platform.YOUTUBE,
                url="https://youtube.com/@pending-export", status=SocialAccountStatus.PENDING,
            ),
            SocialAccount(
                user_id=blogger.id, platform=Platform.VK,
                url="https://vk.com/video/@rejected-export", status=SocialAccountStatus.REJECTED,
            ),
        ])
        db.commit()
        rows = load_data_export_rows(
            db,
            export_type=ExportType.SOCIAL_ACCOUNTS,
            filters=DataExportFilters(date_from=date(2026, 9, 1), date_to=date(2026, 9, 30)),
        )

    assert {row["status"] for row in rows} == {
        SocialAccountStatus.PENDING, SocialAccountStatus.REJECTED,
    }


def test_view_reading_export_keeps_latest_row_per_publication_and_period(client):
    with client.app.state.test_session() as db:
        blogger = User(
            email=f"reading-export-{uuid.uuid4()}@example.com",
            password_hash=hash_password(PASSWORD), role=Role.BLOGGER,
            status=AccountStatus.ACTIVE,
            email_verified_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
        )
        db.add(blogger)
        db.flush()
        account = SocialAccount(
            user_id=blogger.id, platform=Platform.YOUTUBE,
            url=f"https://youtube.com/@reading-{uuid.uuid4()}",
            status=SocialAccountStatus.APPROVED,
        )
        card = VideoCard(
            blogger_id=blogger.id, title="Ролик для выгрузки",
            reported_brand=Brand.AMP, reported_product_name="Тестовый товар",
        )
        db.add_all([account, card])
        db.flush()
        publication = Publication(
            video_card_id=card.id, social_account_id=account.id,
            platform=Platform.YOUTUBE,
            submitted_url=f"https://youtube.com/shorts/{uuid.uuid4().hex[:11]}",
            normalized_url=f"https://youtube.com/shorts/{uuid.uuid4().hex[:11]}",
            external_id=uuid.uuid4().hex[:11], status=PublicationStatus.APPROVED,
            parse_status=PublicationParseStatus.PARSED,
        )
        db.add(publication)
        db.flush()
        period = date(2026, 9, 1)
        db.add_all([
            ViewReading(
                publication_id=publication.id, reporting_period=period,
                source=ReadingSource.YOUTUBE_API, reported_value=100,
                accepted_value=100, status=ReadingStatus.ACCEPTED,
                risk_flags=[], idempotency_key="older",
                captured_at=datetime(2026, 9, 10, tzinfo=timezone.utc),
            ),
            ViewReading(
                publication_id=publication.id, reporting_period=period,
                source=ReadingSource.YOUTUBE_API, reported_value=120,
                accepted_value=None, status=ReadingStatus.PENDING,
                risk_flags=[], idempotency_key="newer",
                captured_at=datetime(2026, 9, 11, tzinfo=timezone.utc),
            ),
        ])
        db.commit()
        rows = load_data_export_rows(
            db, export_type=ExportType.VIEW_READINGS,
            filters=DataExportFilters(date_from=period, date_to=period),
        )
        accepted_rows = load_data_export_rows(
            db, export_type=ExportType.VIEW_READINGS,
            filters=DataExportFilters(date_from=period, date_to=period, status="accepted"),
        )

    assert len(rows) == 1
    assert rows[0]["reported_value"] == 120
    assert rows[0]["status"] == ReadingStatus.PENDING
    assert accepted_rows == []


def test_support_export_keeps_one_row_per_ticket(client):
    with client.app.state.test_session() as db:
        blogger = User(
            email=f"support-export-blogger-{uuid.uuid4()}@example.com",
            password_hash=hash_password(PASSWORD), role=Role.BLOGGER,
            status=AccountStatus.ACTIVE,
            email_verified_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
        )
        manager = User(
            email=f"support-export-manager-{uuid.uuid4()}@example.com",
            password_hash=hash_password(PASSWORD), role=Role.MANAGER,
            status=AccountStatus.ACTIVE,
            email_verified_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
        )
        db.add_all([blogger, manager])
        db.flush()
        ticket = SupportTicket(
            ticket_number=f"TEST-{uuid.uuid4().hex[:12]}", blogger_id=blogger.id,
            category=SupportCategory.GENERAL, subject="Проверка выгрузки",
            status=SupportStatus.IN_PROGRESS, assigned_to_user_id=manager.id,
            creation_idempotency_key=uuid.uuid4(), creation_payload_hash="0" * 64,
            created_at=datetime(2026, 9, 12, tzinfo=timezone.utc),
            last_message_at=datetime(2026, 9, 14, tzinfo=timezone.utc),
        )
        db.add(ticket)
        db.flush()
        db.add_all([
            SupportMessage(
                ticket_id=ticket.id, author_user_id=blogger.id, author_role="blogger",
                body="Первое сообщение", idempotency_key=uuid.uuid4(), payload_hash="1" * 64,
                created_at=datetime(2026, 9, 13, tzinfo=timezone.utc),
            ),
            SupportMessage(
                ticket_id=ticket.id, author_user_id=manager.id, author_role="manager",
                body="Последнее сообщение", idempotency_key=uuid.uuid4(), payload_hash="2" * 64,
                created_at=datetime(2026, 9, 14, tzinfo=timezone.utc),
            ),
        ])
        db.commit()
        rows = load_data_export_rows(
            db, export_type=ExportType.SUPPORT_TICKETS,
            filters=DataExportFilters(date_from=date(2026, 9, 1), date_to=date(2026, 9, 30)),
        )

    assert len(rows) == 1
    assert rows[0]["message_count"] == 2
    assert rows[0]["message_body"] == "Последнее сообщение"
    assert rows[0]["message_author"] == manager.email


def test_every_xlsx_template_is_generated_from_versioned_code(tmp_path):
    for export_type in ExportType:
        for schema_version in (1, 2):
            sheet_name, columns = schema_for(export_type, schema_version)
            assert sheet_name
            assert columns
            destination = tmp_path / f"{export_type.value}-v{schema_version}.xlsx"
            generate_export(
                export_type=export_type,
                export_format=ExportFormat.XLSX,
                schema_version=schema_version,
                rows=[],
                destination=destination,
            )
            assert zipfile.is_zipfile(destination)


def test_current_export_schema_localizes_technical_values(tmp_path):
    from openpyxl import load_workbook

    destination = tmp_path / "readings.xlsx"
    generate_export(
        export_type=ExportType.VIEW_READINGS,
        export_format=ExportFormat.XLSX,
        schema_version=2,
        rows=[{
            "platform": "youtube",
            "source": "youtube_api",
            "status": "accepted",
            "risk_flags": ["unusual_growth"],
        }],
        destination=destination,
    )

    sheet = load_workbook(destination, read_only=True).active
    values = list(sheet.iter_rows(values_only=True))[1]
    _, columns = schema_for(ExportType.VIEW_READINGS, 2)
    row = dict(zip((column.key for column in columns), values, strict=True))
    assert row["platform"] == "YouTube"
    assert row["source"] == "YouTube API"
    assert row["status"] == "Принято"
    assert row["risk_flags"] == "Необычно быстрый рост просмотров"


def test_current_moderation_schema_localizes_legacy_event_names(tmp_path):
    from openpyxl import load_workbook

    destination = tmp_path / "moderation.xlsx"
    generate_export(
        export_type=ExportType.MODERATION_HISTORY,
        export_format=ExportFormat.XLSX,
        schema_version=2,
        rows=[{
            "object_type": "publication",
            "event_type": "publication_approve",
            "from_status": "pending_review",
            "to_status": "approved",
            "changes": {"platform": "youtube"},
        }],
        destination=destination,
    )

    values = list(load_workbook(destination, read_only=True).active.iter_rows(values_only=True))[1]
    _, columns = schema_for(ExportType.MODERATION_HISTORY, 2)
    row = dict(zip((column.key for column in columns), values, strict=True))
    assert row["object_type"] == "Публикация"
    assert row["event_type"] == "Публикация одобрена"
    assert row["from_status"] == "На проверке"
    assert row["to_status"] == "Одобрено"
    assert row["changes"] == '{"Площадка":"YouTube"}'


def test_legacy_export_schema_keeps_original_values(tmp_path):
    destination = tmp_path / "readings.csv"
    generate_export(
        export_type=ExportType.VIEW_READINGS,
        export_format=ExportFormat.CSV,
        schema_version=1,
        rows=[{"platform": "youtube", "source": "youtube_api", "status": "accepted"}],
        destination=destination,
    )
    content = destination.read_text(encoding="utf-8-sig")
    assert "youtube;" in content
    assert "youtube_api;" in content
    assert "accepted;" in content


def test_analyst_templates_exclude_direct_contact_bank_and_ip_fields():
    forbidden = {"phone", "telegram", "sbp_phone", "bank_name", "ip_address"}
    for export_type in (
        ExportType.PUBLICATIONS,
        ExportType.VIEW_READINGS,
        ExportType.ACCRUALS,
    ):
        _, columns = schema_for(export_type, 1)
        assert forbidden.isdisjoint(column.key for column in columns)


def test_legacy_payout_export_routes_do_not_exist(client):
    assert client.get("/api/v1/staff/payout-exports").status_code == 404
    assert client.post("/api/v1/staff/payout-exports", json={}).status_code == 404
