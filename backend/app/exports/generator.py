import csv
import json
import re
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Literal

from app.exports.models import ExportFormat, ExportType
from app.payouts.policy import MOSCOW


@dataclass(frozen=True)
class ColumnSpec:
    key: str
    title: str
    width: int
    kind: Literal["text", "money", "date", "datetime", "integer", "json"] = "text"

BLOGGERS_V1 = (
    ColumnSpec("blogger_id", "ID блогера", 38), ColumnSpec("email", "Email", 32),
    ColumnSpec("full_name", "ФИО", 32), ColumnSpec("display_name", "Отображаемое имя", 28),
    ColumnSpec("phone", "Телефон", 18), ColumnSpec("telegram", "Telegram", 24),
    ColumnSpec("city_country", "Город / страна", 24), ColumnSpec("recipient_status", "Статус получателя", 20),
    ColumnSpec("account_status", "Статус аккаунта", 18), ColumnSpec("profile_status", "Статус профиля", 18),
    ColumnSpec("created_at", "Дата регистрации", 20, "datetime"),
)
SOCIAL_ACCOUNTS_V1 = (
    ColumnSpec("account_id", "ID аккаунта", 38), ColumnSpec("blogger_id", "ID блогера", 38),
    ColumnSpec("blogger", "Блогер", 28), ColumnSpec("platform", "Площадка", 16),
    ColumnSpec("url", "Ссылка", 60), ColumnSpec("follower_count", "Подписчики", 16, "integer"),
    ColumnSpec("status", "Статус", 18), ColumnSpec("updated_at", "Последнее изменение", 20, "datetime"),
)
PUBLICATIONS_V1 = (
    ColumnSpec("publication_id", "ID публикации", 38), ColumnSpec("blogger_id", "ID блогера", 38),
    ColumnSpec("blogger", "Блогер", 28), ColumnSpec("video_card_id", "ID ролика", 38),
    ColumnSpec("title", "Название ролика", 36), ColumnSpec("brand", "Бренд", 18),
    ColumnSpec("product", "Товар", 32), ColumnSpec("platform", "Площадка", 16),
    ColumnSpec("url", "Ссылка", 60), ColumnSpec("external_id", "Внешний ID", 28),
    ColumnSpec("status", "Статус", 20), ColumnSpec("availability", "Доступность", 18),
    ColumnSpec("submitted_at", "Дата отправки", 20, "datetime"), ColumnSpec("reviewed_at", "Дата проверки", 20, "datetime"),
)
VIEW_READINGS_V1 = (
    ColumnSpec("reading_id", "ID показания", 38), ColumnSpec("publication_id", "ID публикации", 38),
    ColumnSpec("blogger_id", "ID блогера", 38), ColumnSpec("platform", "Площадка", 16),
    ColumnSpec("url", "Ссылка", 60), ColumnSpec("period", "Отчётный период", 16, "date"),
    ColumnSpec("source", "Источник", 18), ColumnSpec("reported_value", "Передано просмотров", 22, "integer"),
    ColumnSpec("accepted_value", "Принято просмотров", 22, "integer"), ColumnSpec("status", "Статус", 18),
    ColumnSpec("risk_flags", "Риски", 40, "json"), ColumnSpec("captured_at", "Зафиксировано", 20, "datetime"),
)
MODERATION_HISTORY_V1 = (
    ColumnSpec("event_id", "ID события", 38), ColumnSpec("object_type", "Тип объекта", 22),
    ColumnSpec("object_id", "ID объекта", 38), ColumnSpec("actor_user_id", "ID сотрудника", 38),
    ColumnSpec("event_type", "Событие", 24), ColumnSpec("from_status", "Исходный статус", 20),
    ColumnSpec("to_status", "Новый статус", 20), ColumnSpec("reason", "Причина", 40),
    ColumnSpec("changes", "Изменения", 60, "json"), ColumnSpec("created_at", "Дата события", 20, "datetime"),
)
ACCRUALS_V1 = (
    ColumnSpec("accrual_id", "ID начисления", 38), ColumnSpec("period", "Период", 16, "date"),
    ColumnSpec("blogger_id", "ID блогера", 38), ColumnSpec("blogger", "Блогер", 28),
    ColumnSpec("publication_id", "ID публикации", 38), ColumnSpec("platform", "Площадка", 16),
    ColumnSpec("brand", "Бренд", 18), ColumnSpec("product", "Товар", 32),
    ColumnSpec("eligible_views", "Оплачиваемые просмотры", 24, "integer"), ColumnSpec("rate_kopecks", "Ставка, коп.", 16, "integer"),
    ColumnSpec("amount_kopecks", "Начислено, руб.", 18, "money"), ColumnSpec("adjustment_kopecks", "Корректировка, руб.", 20, "money"),
    ColumnSpec("risk_flags", "Риски", 40, "json"),
)
SUPPORT_TICKETS_V1 = (
    ColumnSpec("ticket_id", "ID обращения", 38), ColumnSpec("ticket_number", "Номер", 24),
    ColumnSpec("blogger_id", "ID блогера", 38), ColumnSpec("blogger", "Блогер", 28),
    ColumnSpec("category", "Категория", 22), ColumnSpec("subject", "Тема", 40), ColumnSpec("status", "Статус", 20),
    ColumnSpec("assigned_to", "Ответственный", 32), ColumnSpec("message_author", "Автор сообщения", 32),
    ColumnSpec("message_body", "Сообщение", 60), ColumnSpec("message_created_at", "Дата сообщения", 20, "datetime"),
    ColumnSpec("created_at", "Создано", 20, "datetime"), ColumnSpec("last_message_at", "Последнее сообщение", 20, "datetime"),
    ColumnSpec("resolved_at", "Решено", 20, "datetime"), ColumnSpec("closed_at", "Закрыто", 20, "datetime"),
)
AUDIT_LOG_V1 = (
    ColumnSpec("event_id", "ID события", 38), ColumnSpec("occurred_at", "Дата события", 20, "datetime"),
    ColumnSpec("actor_user_id", "ID пользователя", 38), ColumnSpec("actor_role", "Роль", 16),
    ColumnSpec("action", "Действие", 34), ColumnSpec("result", "Результат", 16),
    ColumnSpec("object_type", "Тип объекта", 22), ColumnSpec("object_id", "ID объекта", 38),
    ColumnSpec("request_id", "ID запроса", 38), ColumnSpec("ip_address", "IP-адрес", 20),
    ColumnSpec("metadata", "Метаданные", 60, "json"),
)
REGISTER_V1 = (
    ColumnSpec("request_number", "Номер заявки", 32),
    ColumnSpec("recipient_name", "Получатель", 32),
    ColumnSpec("recipient_type", "Тип получателя", 20),
    ColumnSpec("sbp_phone", "Телефон СБП", 18),
    ColumnSpec("bank_name", "Банк", 24),
    ColumnSpec("amount_rubles", "Сумма, руб.", 16, "money"),
    ColumnSpec("requested_on", "Дата заявки", 14, "date"),
    ColumnSpec("approved_on", "Дата одобрения", 16, "date"),
    ColumnSpec("payment_due_date", "Срок оплаты", 14, "date"),
    ColumnSpec("self_employment_verified", "Самозанятость проверена", 24),
    ColumnSpec("manager_comment", "Комментарий", 40),
    ColumnSpec("status", "Статус", 18),
)

HISTORY_V1 = REGISTER_V1 + (
    ColumnSpec("paid_on", "Дата оплаты", 14, "date"),
    ColumnSpec("payment_reference", "Номер операции", 28),
    ColumnSpec("rejected_on", "Дата отказа", 14, "date"),
    ColumnSpec("rejection_reason", "Причина отказа", 40),
    ColumnSpec("receipt_due_date", "Срок чека", 14, "date"),
    ColumnSpec("receipt_received_on", "Дата получения чека", 20, "date"),
    ColumnSpec("approved_by", "Одобрил", 30),
    ColumnSpec("paid_by", "Оплатил", 30),
    ColumnSpec("rejected_by", "Отклонил", 30),
)

_SCHEMAS = {
    (ExportType.BLOGGERS, 1): ("Блогеры", BLOGGERS_V1),
    (ExportType.SOCIAL_ACCOUNTS, 1): ("Социальные аккаунты", SOCIAL_ACCOUNTS_V1),
    (ExportType.PUBLICATIONS, 1): ("Публикации", PUBLICATIONS_V1),
    (ExportType.VIEW_READINGS, 1): ("История просмотров", VIEW_READINGS_V1),
    (ExportType.MODERATION_HISTORY, 1): ("История модерации", MODERATION_HISTORY_V1),
    (ExportType.ACCRUALS, 1): ("Начисления", ACCRUALS_V1),
    (ExportType.PAYOUT_REGISTER, 1): ("Реестр выплат", REGISTER_V1),
    (ExportType.PAYOUT_HISTORY, 1): ("История выплат", HISTORY_V1),
    (ExportType.SUPPORT_TICKETS, 1): ("Обращения", SUPPORT_TICKETS_V1),
    (ExportType.AUDIT_LOG, 1): ("Журнал действий", AUDIT_LOG_V1),
}
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _safe_text(value: Any) -> str:
    if value is None:
        return ""
    value = getattr(value, "value", value)
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    text = _ILLEGAL_EXCEL_CHARACTERS.sub("", str(value))
    return "'" + text if text.lstrip().startswith(_FORMULA_PREFIXES) else text


def moscow_date(value: datetime | date | None) -> date | None:
    if value is None or isinstance(value, date) and not isinstance(value, datetime):
        return value
    if value.tzinfo is None or value.utcoffset() is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(MOSCOW).date()


def moscow_datetime(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None or value.utcoffset() is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(MOSCOW).replace(tzinfo=None)

def _cell_value(row: dict[str, Any], column: ColumnSpec) -> Any:
    value = row.get(column.key)
    if column.kind == "text":
        return _safe_text(value)
    if column.kind == "date":
        return moscow_date(value)
    if column.kind == "datetime":
        return moscow_datetime(value)
    if column.kind == "integer":
        return value
    if column.kind == "json":
        return _safe_text(value)
    if value is None:
        return None
    return (Decimal(value) / Decimal(100)).quantize(Decimal("0.01"))


def schema_for(export_type: ExportType, schema_version: int):
    try:
        return _SCHEMAS[(export_type, schema_version)]
    except KeyError as error:
        raise ValueError("Unsupported export schema version") from error


def generate_export(
    *,
    export_type: ExportType,
    export_format: ExportFormat,
    schema_version: int,
    rows: list[dict[str, Any]],
    destination: Path,
) -> None:
    sheet_name, columns = schema_for(export_type, schema_version)
    if export_format == ExportFormat.CSV:
        _generate_csv(destination, columns, rows)
    else:
        _generate_xlsx(destination, sheet_name, columns, rows)


def _generate_csv(
    destination: Path,
    columns: tuple[ColumnSpec, ...],
    rows: list[dict[str, Any]],
) -> None:
    with destination.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
        writer.writerow(column.title for column in columns)
        for row in rows:
            writer.writerow(
                value.strftime("%d.%m.%Y %H:%M:%S")
                if isinstance(value := _cell_value(row, column), datetime)
                else value.strftime("%d.%m.%Y")
                if isinstance(value, date)
                else format(value, ".2f").replace(".", ",")
                if isinstance(value, Decimal)
                else value
                for column in columns
            )


def _generate_xlsx(
    destination: Path,
    sheet_name: str,
    columns: tuple[ColumnSpec, ...],
    rows: list[dict[str, Any]],
) -> None:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet(sheet_name)
    worksheet.freeze_panes = "A2"
    for index, column in enumerate(columns, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = column.width
    header = []
    for column in columns:
        cell = WriteOnlyCell(worksheet, value=column.title)
        cell.font = Font(bold=True)
        header.append(cell)
    worksheet.append(header)
    for source_row in rows:
        output_row = []
        for column in columns:
            value = _cell_value(source_row, column)
            cell = WriteOnlyCell(worksheet, value=value)
            if column.kind == "money":
                cell.number_format = "0.00"
            elif column.kind == "date":
                cell.number_format = "dd.mm.yyyy"
            elif column.kind == "datetime":
                cell.number_format = "dd.mm.yyyy hh:mm:ss"
            output_row.append(cell)
        worksheet.append(output_row)
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    workbook.save(destination)
